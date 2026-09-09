from __future__ import annotations

import csv
import io
import math
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import load_workbook

from reconcileflow.engine import canonical_company, normalize, normalize_header, number, parse_date


FIELDS = (
    "reference", "date", "company", "service", "pickup", "adult", "child", "infant", "adult_rate", "amount"
)
ALIASES = {
    "reference": {"REFERENCE", "BOOKING REFERENCE", "REFERENCE NO", "BOOKING ID"},
    "date": {"DATE", "SERVICE DATE", "EXECUTION DATE", "TRIP DATE"},
    "company": {"COMPANY", "COMPANY NAME", "AGENCY", "AGENCY NAME"},
    "service": {"SERVICE", "SERVICE NAME", "TRIP", "TRIP TYPE", "EXCURSION"},
    "pickup": {"PICKUP", "PICKUP LOCATION", "HOTEL", "HOTEL NAME", "HOTEL / PICKUP", "MEETING POINT"},
    "adult": {"ADULT", "ADULTS", "ADULT COUNT"},
    "child": {"CHILD", "CHILDREN", "CHILD COUNT"},
    "infant": {"INFANT", "INFANTS", "INFANT COUNT"},
    "adult_rate": {"ADULT RATE", "RATE", "UNIT RATE"},
    "amount": {"AMOUNT", "TOTAL", "TOTAL AMOUNT", "NET AMOUNT"},
}
STATUS_LABELS = {
    "matched": "Matched",
    "mismatch": "Mismatch",
    "missing_a": "Missing in System A",
    "missing_b": "Missing in System B",
    "duplicate": "Duplicate detected",
}
MAX_ROWS = 10_000


def _mapping(headers):
    normalized = [normalize_header(value) for value in headers]
    mapped = {}
    for field, aliases in ALIASES.items():
        candidates = {normalize_header(alias) for alias in aliases}
        for index, header in enumerate(normalized):
            if header in candidates:
                mapped[field] = index
                break
    missing = [field for field in FIELDS if field not in mapped]
    if missing:
        raise ValueError("Missing detailed columns: " + ", ".join(missing) + ".")
    return mapped


def _rows(file_obj: BinaryIO, filename: str):
    suffix = Path(filename).suffix.lower()
    file_obj.seek(0)
    if suffix == ".csv":
        rows = csv.reader(io.StringIO(file_obj.read().decode("utf-8-sig", errors="replace")))
    elif suffix == ".xlsx":
        workbook = load_workbook(file_obj, read_only=True, data_only=True)
        try:
            return [list(row) for index, row in enumerate(workbook.active.iter_rows(values_only=True)) if index <= MAX_ROWS]
        finally:
            workbook.close()
    else:
        raise ValueError("Use an XLSX or CSV file.")
    return [list(row) for index, row in enumerate(rows) if index <= MAX_ROWS]


def finalize_record(raw: dict[str, Any], row_number: int = 0):
    reference = normalize(raw.get("reference")) or "0"
    service_date = parse_date(raw.get("date"))
    company = str(raw.get("company") or "").strip()
    service = str(raw.get("service") or "").strip()
    pickup = str(raw.get("pickup") or "").strip()
    if not service_date:
        raise ValueError(f"Row {row_number}: date is invalid or empty.")
    if not normalize(company) or not normalize(service):
        raise ValueError(f"Row {row_number}: company and service are required.")
    adult = number(raw.get("adult"), "adult", row_number)
    child = number(raw.get("child"), "child", row_number)
    infant = number(raw.get("infant"), "infant", row_number)
    chargeable = adult + child * 0.5
    rate_value = raw.get("adult_rate")
    amount_value = raw.get("amount")
    rate_missing = rate_value in (None, "")
    amount_missing = amount_value in (None, "")
    if rate_missing and amount_missing and chargeable:
        raise ValueError(f"Row {row_number}: provide either adult rate or amount.")
    rate = 0.0 if rate_missing else number(rate_value, "adult rate", row_number)
    amount = 0.0 if amount_missing else number(amount_value, "amount", row_number)
    if rate_missing and chargeable:
        rate = round(amount / chargeable, 4)
    if amount_missing:
        amount = round(rate * chargeable, 2)
    calculated = round(rate * chargeable, 2)
    return {
        "reference": reference,
        "date": service_date,
        "company": company[:200],
        "company_canonical": canonical_company(company),
        "service": service[:200],
        "service_normalized": normalize(service),
        "pickup": pickup[:200],
        "pickup_normalized": normalize(pickup),
        "has_transfer": bool(normalize(pickup)),
        "adult": adult,
        "child": child,
        "infant": infant,
        "chargeable_pax": chargeable,
        "adult_rate": rate,
        "calculated_amount": calculated,
        "pricing_variance": round(amount - calculated, 2),
        "amount": amount,
    }


def read_detailed_dataset(file_obj: BinaryIO, filename: str):
    rows = _rows(file_obj, filename)
    if not rows:
        raise ValueError(f"{filename}: the file is empty.")
    mapped = _mapping(rows[0])
    parsed = []
    for row_number, values in enumerate(rows[1:], start=2):
        if not any(value not in (None, "") for value in values):
            continue
        raw = {field: values[index] if index < len(values) else None for field, index in mapped.items()}
        parsed.append(finalize_record(raw, row_number))
    if not parsed:
        raise ValueError(f"{filename}: no data rows were found.")
    if len(parsed) >= MAX_ROWS:
        raise ValueError(f"The public demo accepts fewer than {MAX_ROWS:,} detailed rows per file.")
    return parsed


def _aggregate(rows):
    grouped = defaultdict(list)
    for row in rows:
        grouped[row["reference"]].append(row)
    result = {}
    for reference, members in grouped.items():
        first = dict(members[0])
        for field in ("adult", "child", "infant", "chargeable_pax", "calculated_amount", "pricing_variance", "amount"):
            first[field] = round(sum(row[field] for row in members), 4 if field == "chargeable_pax" else 2)
        first["adult_rate"] = round(first["amount"] / first["chargeable_pax"], 4) if first["chargeable_pax"] else 0
        first["row_count"] = len(members)
        result[reference] = first
    return result


def reconcile_detailed(system_a, system_b, tolerance=0.01):
    left_rows, right_rows = _aggregate(system_a), _aggregate(system_b)
    output = []
    for reference in sorted(set(left_rows) | set(right_rows)):
        left, right = left_rows.get(reference), right_rows.get(reference)
        differences = []
        if left and right:
            comparisons = (
                ("date", "Service date"), ("company_canonical", "Company"),
                ("service_normalized", "Trip type"), ("pickup_normalized", "Hotel / pickup"),
                ("adult", "Adult count"), ("child", "Child count"), ("infant", "Infant count"),
            )
            differences.extend(label for field, label in comparisons if left[field] != right[field])
            if abs(left["adult_rate"] - right["adult_rate"]) >= tolerance:
                differences.append("Adult rate")
            if abs(left["amount"] - right["amount"]) >= tolerance:
                differences.append("Amount")
            duplicate = left["row_count"] > 1 or right["row_count"] > 1
            status = "duplicate" if duplicate else ("mismatch" if differences else "matched")
        elif left:
            status, differences = "missing_b", ["Record availability"]
        else:
            status, differences = "missing_a", ["Record availability"]
        output.append({
            "reference": reference,
            "status": status,
            "status_label": STATUS_LABELS[status],
            "differences": differences,
            "system_a": left,
            "system_b": right,
            "amount_variance": round((left or {}).get("amount", 0) - (right or {}).get("amount", 0), 2),
        })
    counts = Counter(row["status"] for row in output)
    total = len(output)
    return {
        "summary": {
            "total": total, "matched": counts["matched"], "mismatch": counts["mismatch"],
            "missing_a": counts["missing_a"], "missing_b": counts["missing_b"], "duplicates": counts["duplicate"],
            "match_rate": round(counts["matched"] * 100 / total, 1) if total else 0,
            "amount_variance": round(sum(row["amount_variance"] for row in output), 2),
            "pricing_variance_a": round(sum(row["pricing_variance"] for row in left_rows.values()), 2),
            "pricing_variance_b": round(sum(row["pricing_variance"] for row in right_rows.values()), 2),
        },
        "results": output,
    }
