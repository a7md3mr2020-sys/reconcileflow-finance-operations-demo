from __future__ import annotations

import csv
import io
import math
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import load_workbook

from reconcileflow.daily_matching import daily_matches


FIELDS = ("reference", "date", "company", "service", "adult", "child", "infant", "amount")
MAX_DATA_ROWS = 10_000
HEADER_ALIASES = {
    "reference": {"REFERENCE", "REFERENCENO", "BOOKINGREFERENCE", "REF", "BOOKINGID"},
    "date": {"DATE", "SERVICEDATE", "EXECUTIONDATE", "TRIPDATE"},
    "company": {"COMPANY", "COMPANYNAME", "AGENCY", "AGENCYNAME", "CUSTOMER"},
    "service": {"SERVICE", "SERVICENAME", "TRIP", "TRIPTYPE", "EXCURSION"},
    "adult": {"ADULT", "ADULTS", "ADULTCOUNT"},
    "child": {"CHILD", "CHILDREN", "CHILDCOUNT"},
    "infant": {"INFANT", "INFANTS", "INFANTCOUNT"},
    "amount": {"AMOUNT", "TOTAL", "TOTALAMOUNT", "NETAMOUNT"},
}
COMPANY_ALIASES = {
    "RED SEA TOURS LLC": "RED SEA TOURS",
    "REDSEA TOURS": "RED SEA TOURS",
    "NILE HORIZON TRAVEL CO": "NILE HORIZON TRAVEL",
    "BLUEWAVE TRAVEL": "BLUE WAVE TRAVEL",
}
STATUS_LABELS = {
    "matched": "Matched",
    "mismatch": "Mismatch",
    "missing_a": "Missing in System A",
    "missing_b": "Missing in System B",
    "duplicate": "Duplicate detected",
}


def normalize(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).upper().strip()
    return re.sub(r"\s+", " ", "".join(char if char.isalnum() else " " for char in text)).strip()


def normalize_header(value: Any) -> str:
    return normalize(value).replace(" ", "")


def canonical_company(value: Any) -> str:
    normalized = normalize(value)
    return COMPANY_ALIASES.get(normalized, normalized)


def parse_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(value))).date().isoformat()
        except (OverflowError, ValueError):
            return None
    text = str(value or "").strip().split(" ")[0]
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            continue
    return None


def number(value: Any, field: str, row_number: int) -> float:
    if value in (None, ""):
        return 0.0
    try:
        result = float(str(value).replace(",", "").strip())
    except ValueError as exc:
        raise ValueError(f"Row {row_number}: {field} must be numeric.") from exc
    if not math.isfinite(result) or result < 0:
        raise ValueError(f"Row {row_number}: {field} must be zero or a positive number.")
    return result


def _mapping(headers: list[Any]) -> dict[str, int]:
    normalized_headers = [normalize_header(value) for value in headers]
    result: dict[str, int] = {}
    for field, aliases in HEADER_ALIASES.items():
        normalized_aliases = {normalize_header(value) for value in aliases}
        for index, header in enumerate(normalized_headers):
            if header in normalized_aliases:
                result[field] = index
                break
    missing = [field for field in FIELDS if field not in result]
    if missing:
        raise ValueError("Missing required columns: " + ", ".join(missing) + ".")
    return result


def _raw_rows(file_obj: BinaryIO, filename: str) -> list[list[Any]]:
    suffix = Path(filename).suffix.lower()
    file_obj.seek(0)
    if suffix == ".csv":
        raw = file_obj.read()
        text = raw.decode("utf-8-sig", errors="replace")
        rows = []
        for index, row in enumerate(csv.reader(io.StringIO(text))):
            if index > MAX_DATA_ROWS:
                raise ValueError(f"The public demo accepts up to {MAX_DATA_ROWS:,} rows per file.")
            rows.append(list(row))
        return rows
    if suffix == ".xlsx":
        workbook = load_workbook(file_obj, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index > MAX_DATA_ROWS:
                    raise ValueError(f"The public demo accepts up to {MAX_DATA_ROWS:,} rows per file.")
                rows.append(list(row))
            return rows
        finally:
            workbook.close()
    raise ValueError("Use an XLSX or CSV file.")


def read_dataset(file_obj: BinaryIO, filename: str) -> list[dict[str, Any]]:
    rows = _raw_rows(file_obj, filename)
    if not rows:
        raise ValueError(f"{filename}: the file is empty.")
    mapping = _mapping(rows[0])
    parsed: list[dict[str, Any]] = []
    for row_number, values in enumerate(rows[1:], start=2):
        if not any(value not in (None, "") for value in values):
            continue
        get = lambda field: values[mapping[field]] if mapping[field] < len(values) else None
        reference = normalize(get("reference")) or "0"
        service_date = parse_date(get("date"))
        company = str(get("company") or "").strip()
        service = str(get("service") or "").strip()
        if not service_date:
            raise ValueError(f"Row {row_number}: date is invalid or empty.")
        if not normalize(company):
            raise ValueError(f"Row {row_number}: company is required.")
        if not normalize(service):
            raise ValueError(f"Row {row_number}: service is required.")
        parsed.append({
            "reference": reference,
            "date": service_date,
            "company": company[:200],
            "company_canonical": canonical_company(company),
            "service": service[:200],
            "service_normalized": normalize(service),
            "adult": number(get("adult"), "adult", row_number),
            "child": number(get("child"), "child", row_number),
            "infant": number(get("infant"), "infant", row_number),
            "amount": number(get("amount"), "amount", row_number),
        })
    if not parsed:
        raise ValueError(f"{filename}: no data rows were found.")
    return parsed


def _aggregate(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[row["reference"]].append(row)
    result = {}
    for reference, members in grouped.items():
        first = members[0]
        result[reference] = {
            **first,
            "adult": sum(row["adult"] for row in members),
            "child": sum(row["child"] for row in members),
            "infant": sum(row["infant"] for row in members),
            "amount": round(sum(row["amount"] for row in members), 2),
            "row_count": len(members),
        }
    return result


def reconcile(
    system_a: list[dict[str, Any]],
    system_b: list[dict[str, Any]],
    amount_tolerance: float = 0.01,
) -> dict[str, Any]:
    a_rows = _aggregate(system_a)
    b_rows = _aggregate(system_b)
    results = []
    mappings = {}
    for reference in sorted(set(a_rows) | set(b_rows)):
        left = a_rows.get(reference)
        right = b_rows.get(reference)
        differences: list[str] = []
        if left and right:
            mappings[(left["company"], right["company"])] = {
                "system_a": left["company"],
                "system_b": right["company"],
                "system_a_canonical": left["company_canonical"],
                "system_b_canonical": right["company_canonical"],
                "mapped": left["company_canonical"] == right["company_canonical"],
            }
            comparisons = (
                ("date", "Service date"),
                ("company_canonical", "Company"),
                ("service_normalized", "Service"),
                ("adult", "Adult count"),
                ("child", "Child count"),
                ("infant", "Infant count"),
            )
            differences.extend(label for field, label in comparisons if left[field] != right[field])
            if abs(left["amount"] - right["amount"]) >= amount_tolerance:
                differences.append("Amount")
            duplicate = left["row_count"] > 1 or right["row_count"] > 1
            status = "duplicate" if duplicate else ("mismatch" if differences else "matched")
        elif left:
            status, differences = "missing_b", ["Record availability"]
        else:
            status, differences = "missing_a", ["Record availability"]
        results.append({
            "reference": reference,
            "status": status,
            "status_label": STATUS_LABELS[status],
            "differences": differences,
            "system_a": left,
            "system_b": right,
            "amount_variance": round((left["amount"] if left else 0) - (right["amount"] if right else 0), 2),
        })
    counts = Counter(row["status"] for row in results)
    total = len(results)
    matched = counts["matched"]
    return {
        "summary": {
            "total": total,
            "matched": matched,
            "mismatch": counts["mismatch"],
            "missing_a": counts["missing_a"],
            "missing_b": counts["missing_b"],
            "duplicates": counts["duplicate"],
            "match_rate": round(matched * 100 / total, 1) if total else 0,
            "system_a_rows": len(system_a),
            "system_b_rows": len(system_b),
            "amount_variance": round(sum(row["amount_variance"] for row in results), 2),
        },
        "results": results,
        "source_rows_a": system_a,
        "source_rows_b": system_b,
        "daily_results": daily_matches(system_a, system_b),
        "mappings": sorted(mappings.values(), key=lambda row: (row["system_a_canonical"], row["system_a"])),
    }
