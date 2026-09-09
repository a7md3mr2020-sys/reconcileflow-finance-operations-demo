from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from reconcileflow.sample_data import HEADERS, SYSTEM_A, SYSTEM_B
from reconcileflow.extended_sample_data import DETAILED_A, DETAILED_B, DETAILED_HEADERS


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "sample-data"


def write_workbook(name, title, headers, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173E47")
        cell.alignment = Alignment(horizontal="center")
    widths = [18, 15, 28, 24, 26, 12, 12, 12, 16, 16]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + column)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(OUTPUT / name)


if __name__ == "__main__":
    OUTPUT.mkdir(exist_ok=True)
    write_workbook("system-a-demo.xlsx", "System A", HEADERS, SYSTEM_A)
    write_workbook("system-b-demo.xlsx", "System B", HEADERS, SYSTEM_B)
    write_workbook("detailed-system-a-demo.xlsx", "Detailed A", DETAILED_HEADERS, DETAILED_A)
    write_workbook("detailed-system-b-demo.xlsx", "Detailed B", DETAILED_HEADERS, DETAILED_B)
    print(f"Created synthetic sample workbooks in {OUTPUT}")
