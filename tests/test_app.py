import io
import os
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

import app as app_module
from reconcileflow.sample_data import HEADERS, SYSTEM_A, SYSTEM_B


def test_empty_host_environment_values_use_safe_defaults():
    environment = os.environ.copy()
    environment.update({"DEMO_DB_PATH": "", "RUN_RETENTION_HOURS": ""})
    result = subprocess.run(
        [sys.executable, "-c", "import app; print(app.RUN_RETENTION_HOURS, app.DATABASE_PATH.name)"],
        cwd=Path(__file__).resolve().parents[1],
        env=environment,
        capture_output=True,
        text=True,
        timeout=30,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    assert result.stdout.strip() == "24 reconcileflow-demo.db"


def csrf(client):
    with client.session_transaction() as session:
        session["csrf_token"] = "test-token"
    return "test-token"


def sample_id(client, mode, month):
    with client.session_transaction() as state:
        return f"portfolio-{mode}-{month}-{state['visitor_id']}"


def workbook_bytes(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def test_health_and_security_headers(client):
    response = client.get("/health")
    assert response.json == {"ok": True, "service": "ReconcileFlow"}
    assert response.headers["X-Content-Type-Options"] == "nosniff"
    assert response.headers["X-Frame-Options"] == "DENY"


def test_hub_exposes_all_three_workflows_and_report_workshop(client):
    response = client.get("/")
    assert response.status_code == 200
    for label in (b"Standard Reconciliation", b"Detailed Reconciliation", b"Invoice Tracking", b"Report Workshop"):
        assert label in response.data
    assert b"Try the live standard demo" in response.data
    assert b"Try the live detailed demo" in response.data


def test_hr_live_demo_links_open_ready_combined_workspaces(client):
    standard = client.get("/try/standard", follow_redirects=True)
    assert standard.status_code == 200
    assert b"Projects operating as one reconciliation workspace" in standard.data
    assert b"May 2026 marine operations" in standard.data
    assert b"June 2026 marine operations" in standard.data
    detailed = client.get("/try/detailed", follow_redirects=True)
    assert detailed.status_code == 200
    assert b"Hotel or pickup A / B" in detailed.data
    assert client.get("/try/unsupported").status_code == 404


def test_complete_demo_workflow_and_export(client):
    response = client.post("/demo", data={"csrf_token": csrf(client)}, follow_redirects=False)
    assert response.status_code == 302
    result_url = response.headers["Location"]
    page = client.get(result_url)
    assert page.status_code == 200
    assert b"13" in page.data
    assert b"Duplicate detected" in page.data
    export = client.get(result_url + "/export.xlsx")
    assert export.status_code == 200
    workbook = load_workbook(io.BytesIO(export.data), read_only=True, data_only=True)
    assert workbook.sheetnames == ["Summary", "Reconciliation Results", "Daily Matching", "Company Mappings"]
    assert workbook["Daily Matching"].max_row > 1
    assert workbook["Reconciliation Results"].max_row == 14
    assert workbook["Reconciliation Results"]["A2"].value == "RF-1001"
    assert workbook["Reconciliation Results"].max_column == 20
    workbook.close()


def test_upload_validation_returns_friendly_error(client):
    response = client.post(
        "/reconcile",
        data={
            "csrf_token": csrf(client),
            "project_name": "Invalid upload",
            "system_a": (io.BytesIO(b"not excel"), "a.xlsx"),
            "system_b": (io.BytesIO(b"not excel"), "b.xlsx"),
        },
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert b"invalid Excel file" in response.data


def test_valid_excel_upload_runs_reconciliation(client):
    response = client.post(
        "/reconcile",
        data={
            "csrf_token": csrf(client),
            "project_name": "Uploaded synthetic review",
            "system_a": (workbook_bytes(SYSTEM_A), "a.xlsx"),
            "system_b": (workbook_bytes(SYSTEM_B), "b.xlsx"),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b"Uploaded synthetic review" in response.data
    assert b"15.4%" in response.data


def test_sample_downloads_and_excel_cell_protection(client):
    response = client.get("/sample-data/a.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data), read_only=True, data_only=True)
    assert workbook.active.max_row == 14
    workbook.close()
    assert app_module.excel_safe("=2+2") == "'=2+2"
    assert app_module.excel_safe("Normal label") == "Normal label"


def test_csrf_is_required(client):
    response = client.post("/demo", data={})
    assert response.status_code == 400


def test_csrf_failure_returns_json_for_interactive_requests(client):
    response = client.post(
        "/projects/missing/adjustment",
        data={"reference": "RF-1001", "side": "system_a", "amount": "10"},
        headers={"Accept": "application/json"},
    )

    assert response.status_code == 400
    assert response.is_json
    assert response.get_json() == {
        "ok": False,
        "error": "The form expired. Refresh the page and try again.",
    }


def test_unknown_result_is_404(client):
    assert client.get("/results/unknown").status_code == 404


def test_language_switch_persists_and_rejects_open_redirects(client):
    response = client.get("/language/ar?next=/detailed", follow_redirects=False)
    assert response.status_code == 302
    assert response.headers["Location"].endswith("/detailed")
    page = client.get("/detailed")
    assert b'<html lang="ar" dir="rtl">' in page.data
    assert b'/static/i18n.js' in page.data

    unsafe = client.get("/language/en?next=https://example.com", follow_redirects=False)
    assert unsafe.headers["Location"].endswith("/")
    unsafe_backslash = client.get("/language/en?next=/\\example.com", follow_redirects=False)
    assert unsafe_backslash.headers["Location"].endswith("/")
    assert client.get("/language/fr").status_code == 404


def test_detailed_demo_and_export(client):
    response = client.post("/detailed/demo", data={"csrf_token": csrf(client)}, follow_redirects=False)
    assert response.status_code == 302
    page = client.get(response.headers["Location"])
    assert page.status_code == 200
    assert b"Reference-number reconciliation" in page.data
    assert b"Hotel or pickup" in page.data
    export = client.get(response.headers["Location"] + "/export.xlsx")
    workbook = load_workbook(io.BytesIO(export.data), read_only=True, data_only=True)
    assert workbook.sheetnames == ["Summary", "Detailed Results", "Daily Matching", "Reference Matching"]
    assert workbook["Daily Matching"].max_row > 1
    daily_headers = [cell.value for cell in workbook["Daily Matching"][1]]
    assert "Calculated Variance" in daily_headers
    assert "Pickup A" in daily_headers
    assert "Adult Rate B" in daily_headers
    assert workbook["Detailed Results"].max_row == 8
    workbook.close()


def test_invoice_tracking_views_export_and_workshop(client):
    for tab in ("overview", "companies", "names", "analysis", "activity"):
        response = client.get(f"/invoice-tracking?tab={tab}")
        assert response.status_code == 200
        assert b"Marine operations invoice follow-up" in response.data
    export = client.get("/invoice-tracking/export.xlsx")
    workbook = load_workbook(io.BytesIO(export.data), read_only=True, data_only=True)
    assert workbook.sheetnames == ["Summary", "Company Tracking", "Activity Log"]
    assert workbook["Company Tracking"].max_row == 8
    workbook.close()
    workshop = client.get("/report-workshop")
    assert workshop.status_code == 200
    assert b"Visible identity" in workshop.data
    assert b"Authorized signature and stamp" in workshop.data


def test_invoice_updates_flow_through_balances_activity_export_and_report(client):
    assert client.get("/invoice-tracking?tab=names").status_code == 200
    token = csrf(client)
    payment = client.post(
        "/invoice-tracking/company/CMP-001/tracking",
        data={"csrf_token": token, "payment": "10000.00", "stage": "paid"},
    )
    assert payment.status_code == 200
    assert payment.json["company"]["balance"] == 25000
    assert payment.json["company"]["progress"] == 100
    assert payment.json["summary"]["collected"] == 335000

    identity = client.post(
        "/invoice-tracking/company/CMP-001/identity",
        data={"csrf_token": token, "name": "Nile Horizon Tours", "invoice": "130000.00", "paid": "95000.00", "notes": "Reviewed"},
    )
    assert identity.status_code == 200
    assert identity.json["company"]["balance"] == 25000
    assert identity.json["summary"]["invoice"] == 477500
    assert identity.json["summary"]["collected"] == 340000
    assert b"Nile Horizon Tours" in client.get("/invoice-tracking?tab=companies").data

    added = client.post(
        "/invoice-tracking/companies",
        data={"csrf_token": token, "name": "Harbor Finance Travel", "invoice": "1000.00", "paid": "200.00"},
    )
    assert added.status_code == 200
    assert added.json["company"]["code"] == "CMP-008"
    assert added.json["company"]["balance"] == 800
    assert added.json["summary"]["companies"] == 8
    assert b"Harbor Finance Travel" in client.get("/report-workshop").data

    activity = client.get("/invoice-tracking?tab=activity")
    assert b"Payment or stage updated" in activity.data
    assert b"Company identity or opening amounts updated" in activity.data
    export = client.get("/invoice-tracking/export.xlsx")
    workbook = load_workbook(io.BytesIO(export.data), read_only=True, data_only=True)
    assert workbook["Company Tracking"].max_row == 9
    assert workbook["Activity Log"].max_row == 4
    assert workbook["Summary"]["B2"].value == 8
    assert workbook["Summary"]["B3"].value == 478500
    workbook.close()

    second = app_module.app.test_client()
    assert b"Harbor Finance Travel" not in second.get("/invoice-tracking?tab=names").data
    assert second.post(
        "/invoice-tracking/company/CMP-008/tracking",
        data={"csrf_token": csrf(second), "payment": "1.00", "stage": "paid"},
    ).status_code == 404


def test_invoice_rejects_invalid_amounts_and_keeps_financial_state(client):
    client.get("/invoice-tracking")
    token = csrf(client)
    for value in ("-1", "nan", "1000000000", "wrong"):
        response = client.post(
            "/invoice-tracking/company/CMP-001/tracking",
            data={"csrf_token": token, "payment": value, "stage": "sent"},
        )
        assert response.status_code == 400
    assert client.post(
        "/invoice-tracking/company/CMP-001/identity",
        data={"csrf_token": token, "name": "", "invoice": "1", "paid": "1"},
    ).status_code == 400
    with app_module.app.test_request_context():
        with client.session_transaction() as state:
            app_module.session.update(state)
        assert app_module.invoice_demo_payload()["companies"][0]["balance"] == 30000


def test_invoice_report_and_excel_use_complete_filtered_company_set(client):
    export = client.get("/invoice-tracking/export.xlsx?direction=payable")
    assert export.status_code == 200
    workbook = load_workbook(io.BytesIO(export.data), read_only=True, data_only=True)
    assert workbook["Summary"]["B2"].value == 1
    assert workbook["Summary"]["B3"].value == 72000
    assert workbook["Summary"]["B6"].value == 8000
    assert workbook["Summary"]["B7"].value == -8000
    assert workbook["Company Tracking"].max_row == 2
    assert workbook["Company Tracking"]["B2"].value == "Blue Wave Travel"
    workbook.close()

    report = client.get("/report-workshop?direction=payable")
    assert report.status_code == 200
    assert b"Blue Wave Travel" in report.data
    assert b"Nile Horizon Travel" not in report.data
    assert b"72,000.00" in report.data

    searched = client.get("/invoice-tracking/export.xlsx?q=coastal")
    workbook = load_workbook(io.BytesIO(searched.data), read_only=True, data_only=True)
    assert workbook["Company Tracking"].max_row == 2
    assert workbook["Company Tracking"]["B2"].value == "Coastal Journey Group"
    workbook.close()


def test_operations_control_center_views_and_export(client):
    expected = {
        "overview": b"Four periods acting as one workspace",
        "companies": b"Unified company groups across projects",
        "settlements": b"Settlement movement ledger",
        "no-shows": b"Whole and partial booking evidence",
        "delivery": b"Delivery and approved statement readiness",
    }
    for tab, marker in expected.items():
        response = client.get(f"/control-center?tab={tab}")
        assert response.status_code == 200
        assert marker in response.data

    export = client.get("/control-center/export.xlsx")
    assert export.status_code == 200
    workbook = load_workbook(io.BytesIO(export.data), read_only=True, data_only=True)
    assert workbook.sheetnames == [
        "Executive Summary", "Projects", "Unified Companies", "Settlement Ledger",
        "No Shows", "Delivery Control",
    ]
    assert workbook["Unified Companies"].max_row == 6
    assert workbook["Settlement Ledger"].max_row == 6
    assert workbook["No Shows"]["H5"].value == 210
    workbook.close()


def test_standard_projects_combine_and_adjustments_are_audited(client):
    library = client.get("/standard")
    assert library.status_code == 200
    assert b"Monthly reconciliation projects" in library.data
    may_id = sample_id(client, "standard", "may")
    june_id = sample_id(client, "standard", "june")
    assert may_id.encode() in library.data

    combined = client.get(
        f"/projects/standard/combined?project={may_id}&project={june_id}"
    )
    assert combined.status_code == 200
    assert b"Projects operating as one reconciliation workspace" in combined.data
    assert b"May 2026 marine operations" in combined.data
    assert b"June 2026 marine operations" in combined.data
    sample_daily_count = len(app_module.reconcile(app_module.records(SYSTEM_A), app_module.records(SYSTEM_B))["daily_results"])
    assert combined.data.count(b'data-project-id="combined"') == sample_daily_count

    first = client.post(
        f"/projects/{may_id}/adjustment",
        data={"csrf_token": csrf(client), "project": [may_id, june_id], "reference": "RF-1001", "side": "system_a", "amount": "200"},
    )
    assert first.status_code == 200
    assert first.json["movement"]["before"] == 250
    assert first.json["movement"]["after"] == 200
    assert first.json["movement"]["movement"] == 50
    assert first.json["row"]["amount_variance"] == -50
    assert first.json["summary"]["matched"] == 1
    assert first.json["summary"]["amount_variance"] == 100
    assert "-25.00" in first.json["daily_html"]
    assert "May 2026 marine operations, June 2026 marine operations" in first.json["daily_html"]

    second = client.post(
        f"/projects/{may_id}/adjustment",
        data={"csrf_token": csrf(client), "project": [may_id, june_id], "reference": "RF-1001", "side": "system_a", "amount": "230"},
    )
    assert second.status_code == 200
    assert second.json["movement"]["movement"] == -30
    assert second.json["summary"]["amount_variance"] == 130
    with app_module.database() as connection:
        total = connection.execute(
            "SELECT ROUND(SUM(movement), 2) AS total FROM financial_movements WHERE run_id=?",
            (may_id,),
        ).fetchone()["total"]
        count = connection.execute(
            "SELECT COUNT(*) AS count FROM financial_movements WHERE run_id=?",
            (may_id,),
        ).fetchone()["count"]
    assert total == 20
    assert count == 2
    missing_source = client.post(
        f"/projects/{may_id}/adjustment",
        data={"csrf_token": csrf(client), "reference": "RF-1007", "side": "system_b", "amount": "10"},
    )
    assert missing_source.status_code == 400
    assert len(app_module.movement_rows([may_id])) == 2
    with app_module.app.test_request_context():
        with client.session_transaction() as state:
            app_module.session.update(state)
        run = app_module.get_run(may_id)
    adjusted = next(row for row in run["results"] if row["reference"] == "RF-1001")
    assert adjusted["system_a"]["amount"] == 230
    assert adjusted["amount_variance"] == -20
    assert round(sum(row["variance"] for row in run["daily_results"]), 2) == run["summary"]["amount_variance"]
    export = client.get(
        f"/projects/standard/combined/export.xlsx?project={may_id}&project={june_id}"
    )
    assert export.status_code == 200
    workbook = load_workbook(io.BytesIO(export.data), read_only=True, data_only=True)
    assert workbook.sheetnames == ["Combined Summary", "Combined Bookings", "Daily Matching", "Financial Movements"]
    assert workbook["Combined Bookings"].max_row == 27
    assert workbook["Daily Matching"].max_row > 1
    assert round(sum(row[12] for row in workbook["Daily Matching"].iter_rows(min_row=2, values_only=True)), 2) == workbook["Combined Summary"]["B5"].value
    assert workbook["Financial Movements"].max_row == 3
    assert workbook["Combined Summary"]["B6"].value == 20
    workbook.close()


def test_detailed_projects_adjust_pricing_and_reject_invalid_edits(client):
    assert client.get("/detailed").status_code == 200
    may_id = sample_id(client, "detailed", "may")
    june_id = sample_id(client, "detailed", "june")
    combined = client.get(f"/projects/detailed/combined?project={may_id}&project={june_id}")
    assert combined.status_code == 200
    assert b"Calculated amount A / B" in combined.data
    assert b"May 2026 marine operations, June 2026 marine operations" in combined.data
    with app_module.app.test_request_context():
        with client.session_transaction() as state:
            app_module.session.update(state)
        run = app_module.get_run(may_id)
    row = next(item for item in run["results"] if item.get("system_a"))
    reference = row["reference"]
    source = row["system_a"]
    target = source["amount"] + 15
    response = client.post(
        f"/projects/{may_id}/adjustment",
        data={"csrf_token": csrf(client), "reference": reference, "side": "system_a", "amount": str(target)},
    )
    assert response.status_code == 200
    with app_module.app.test_request_context():
        with client.session_transaction() as state:
            app_module.session.update(state)
        updated = app_module.get_run(may_id)
    updated_row = next(item for item in updated["results"] if item["reference"] == reference)
    assert updated_row["system_a"]["pricing_variance"] == round(
        target - updated_row["system_a"]["calculated_amount"], 2
    )
    assert updated["summary"]["pricing_variance_a"] == round(sum(
        (item.get("system_a") or {}).get("pricing_variance", 0) for item in updated["results"]
    ), 2)

    movement_count = len(app_module.movement_rows([may_id]))
    unchanged = client.post(
        f"/projects/{may_id}/adjustment",
        data={"csrf_token": csrf(client), "reference": reference, "side": "system_a", "amount": str(target)},
    )
    assert unchanged.status_code == 400
    assert len(app_module.movement_rows([may_id])) == movement_count
    invalid = client.post(
        f"/projects/{may_id}/adjustment",
        data={"csrf_token": csrf(client), "reference": reference, "side": "system_a", "amount": "-1"},
    )
    assert invalid.status_code == 400


def test_public_demo_projects_and_adjustments_are_isolated_between_visitors(client):
    first = client
    second = app_module.app.test_client()
    assert first.get("/standard").status_code == 200
    first_may = sample_id(first, "standard", "may")
    assert second.get("/standard").status_code == 200
    second_may = sample_id(second, "standard", "may")
    assert first_may != second_may

    response = first.post(
        f"/projects/{first_may}/adjustment",
        data={"csrf_token": csrf(first), "reference": "RF-1001", "side": "system_a", "amount": "200"},
    )
    assert response.status_code == 200
    assert second.get(f"/projects/standard/combined?project={first_may}").status_code == 302
    assert second.get(f"/projects/standard/combined/export.xlsx?project={first_may}").status_code == 404
    assert second.get(f"/results/{first_may}").status_code == 404
    assert second.post(
        f"/projects/{first_may}/adjustment",
        data={"csrf_token": csrf(second), "reference": "RF-1001", "side": "system_a", "amount": "190"},
    ).status_code == 404

    second_workspace = second.get(f"/projects/standard/combined?project={second_may}")
    assert second_workspace.status_code == 200
    with app_module.app.test_request_context():
        with second.session_transaction() as state:
            app_module.session.update(state)
        second_run = app_module.get_run(second_may)
    row = next(row for row in second_run["results"] if row["reference"] == "RF-1001")
    assert row["system_a"]["amount"] == 250
    assert app_module.movement_rows([second_may]) == []


def test_legacy_sample_links_open_in_current_visitors_workspace(client):
    for mode in ("standard", "detailed"):
        legacy = f"/projects/{mode}/combined?project=portfolio-{mode}-may&project=portfolio-{mode}-june"
        page = client.get(legacy)
        assert page.status_code == 200
        assert b"May 2026 marine operations" in page.data
        assert b"June 2026 marine operations" in page.data
        current_may = sample_id(client, mode, "may")
        assert current_may.encode() in page.data
        assert client.get(legacy.replace("/combined?", "/combined/export.xlsx?")).status_code == 200

        duplicate = client.get(f"/projects/{mode}/combined?project=portfolio-{mode}-may&project={current_may}")
        assert b"<span>Open projects</span><strong>1</strong>" in duplicate.data
        expected_rows = 13 if mode == "standard" else 7
        assert f"data-visible-count>{expected_rows}</strong>".encode() in duplicate.data


import pytest


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(app_module, "DATABASE_PATH", tmp_path / "test.db")
    app_module.init_database()
    app_module.app.config.update(TESTING=True, SECRET_KEY="test-key")
    with app_module.app.test_client() as test_client:
        yield test_client
