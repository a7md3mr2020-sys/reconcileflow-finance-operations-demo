from __future__ import annotations

import io
import json
import math
import os
import secrets
import sqlite3
import tempfile
import zipfile
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path
from uuid import uuid4

from flask import Flask, abort, flash, jsonify, redirect, render_template, request, send_file, session, url_for
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from werkzeug.middleware.proxy_fix import ProxyFix

from reconcileflow.engine import read_dataset, reconcile
from reconcileflow.detailed import read_detailed_dataset, reconcile_detailed
from reconcileflow.invoice_demo import STAGES, amount_to_minor, invoice_view
from reconcileflow.daily_matching import daily_matches
from reconcileflow.control_center import control_center_payload
from reconcileflow.extended_sample_data import (
    DETAILED_A,
    DETAILED_B,
    DETAILED_HEADERS,
    INVOICE_COMPANIES,
    detailed_records,
)
from reconcileflow.sample_data import HEADERS, SYSTEM_A, SYSTEM_B, records


BASE_DIR = Path(__file__).resolve().parent
DATABASE_PATH = Path(os.environ.get("DEMO_DB_PATH") or Path(tempfile.gettempdir()) / "reconcileflow-demo.db")
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
RUN_RETENTION_HOURS = int(os.environ.get("RUN_RETENTION_HOURS") or "24")

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY") or secrets.token_hex(32)
app.config.update(
    MAX_CONTENT_LENGTH=MAX_UPLOAD_BYTES,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.environ.get("COOKIE_SECURE", "0") == "1",
)
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)


def database():
    DATABASE_PATH.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row
    return connection


def init_database():
    with database() as connection:
        connection.execute(
            """CREATE TABLE IF NOT EXISTS demo_runs(
                   id TEXT PRIMARY KEY,
                   project_name TEXT NOT NULL,
                   payload_json TEXT NOT NULL,
                   mode TEXT NOT NULL DEFAULT 'standard',
                   owner_id TEXT NOT NULL DEFAULT '',
                   created_at TEXT NOT NULL
               )"""
        )
        columns = {row["name"] for row in connection.execute("PRAGMA table_info(demo_runs)")}
        if "mode" not in columns:
            connection.execute("ALTER TABLE demo_runs ADD COLUMN mode TEXT NOT NULL DEFAULT 'standard'")
        if "owner_id" not in columns:
            connection.execute("ALTER TABLE demo_runs ADD COLUMN owner_id TEXT NOT NULL DEFAULT ''")
        connection.execute("CREATE INDEX IF NOT EXISTS idx_demo_runs_owner ON demo_runs(owner_id, mode)")
        connection.execute(
            "UPDATE demo_runs SET mode='detailed' WHERE mode='standard' AND LOWER(project_name) LIKE '%detailed%'"
        )
        connection.execute(
            """CREATE TABLE IF NOT EXISTS financial_movements(
                   id INTEGER PRIMARY KEY AUTOINCREMENT,
                   run_id TEXT NOT NULL,
                   mode TEXT NOT NULL,
                   occurred_at TEXT NOT NULL,
                   reference TEXT NOT NULL,
                   side TEXT NOT NULL,
                   before_amount REAL NOT NULL,
                   after_amount REAL NOT NULL,
                   movement REAL NOT NULL,
                   action TEXT NOT NULL
               )"""
        )
        connection.execute(
            "CREATE INDEX IF NOT EXISTS idx_financial_movements_run ON financial_movements(run_id, id DESC)"
        )
        connection.execute(
            """CREATE TABLE IF NOT EXISTS invoice_demo_companies(
                   owner_id TEXT NOT NULL,
                   code TEXT NOT NULL,
                   name TEXT NOT NULL,
                   invoice_minor INTEGER NOT NULL,
                   paid_minor INTEGER NOT NULL,
                   payment_minor INTEGER NOT NULL,
                   stage TEXT NOT NULL,
                   notes TEXT NOT NULL,
                   created_at TEXT NOT NULL,
                   PRIMARY KEY(owner_id, code)
               )"""
        )
        connection.execute(
            """CREATE TABLE IF NOT EXISTS invoice_demo_events(
                   id INTEGER PRIMARY KEY AUTOINCREMENT,
                   owner_id TEXT NOT NULL,
                   company_code TEXT NOT NULL,
                   company_name TEXT NOT NULL,
                   occurred_at TEXT NOT NULL,
                   action TEXT NOT NULL,
                   before_minor INTEGER NOT NULL,
                   after_minor INTEGER NOT NULL
               )"""
        )
        connection.execute(
            "CREATE INDEX IF NOT EXISTS idx_invoice_demo_events_owner ON invoice_demo_events(owner_id, id DESC)"
        )


def visitor_id():
    if "visitor_id" not in session:
        session["visitor_id"] = uuid4().hex
    return session["visitor_id"]


def sample_project_id(mode, month):
    return f"portfolio-{mode}-{month}-{visitor_id()}"


def ensure_invoice_demo():
    owner = visitor_id()
    cutoff = (datetime.now(timezone.utc) - timedelta(hours=RUN_RETENTION_HOURS)).isoformat()
    with database() as connection:
        connection.execute("BEGIN IMMEDIATE")
        expired = [row["owner_id"] for row in connection.execute(
            "SELECT DISTINCT owner_id FROM invoice_demo_companies WHERE created_at < ?", (cutoff,)
        )]
        for old_owner in expired:
            connection.execute("DELETE FROM invoice_demo_events WHERE owner_id=?", (old_owner,))
            connection.execute("DELETE FROM invoice_demo_companies WHERE owner_id=?", (old_owner,))
        exists = connection.execute(
            "SELECT 1 FROM invoice_demo_companies WHERE owner_id=? LIMIT 1", (owner,)
        ).fetchone()
        if not exists:
            created_at = datetime.now(timezone.utc).isoformat()
            connection.executemany(
                """INSERT INTO invoice_demo_companies(
                       owner_id,code,name,invoice_minor,paid_minor,payment_minor,stage,notes,created_at
                   ) VALUES(?,?,?,?,?,?,?,?,?)""",
                [(owner, row["code"], row["name"], amount_to_minor(row["invoice"]),
                  amount_to_minor(row["paid"]), amount_to_minor(row["payment"]),
                  row["stage"], row["notes"], created_at) for row in INVOICE_COMPANIES],
            )


def invoice_demo_payload(search="", direction=""):
    ensure_invoice_demo()
    with database() as connection:
        companies = connection.execute(
            "SELECT * FROM invoice_demo_companies WHERE owner_id=? ORDER BY code", (visitor_id(),)
        ).fetchall()
        events = connection.execute(
            "SELECT * FROM invoice_demo_events WHERE owner_id=? ORDER BY id DESC", (visitor_id(),)
        ).fetchall()
    return invoice_view(companies, events, search=search, direction=direction)


def store_run(project_name, payload, mode="standard", run_id=None):
    if mode not in {"standard", "detailed"}:
        raise ValueError("Unsupported reconciliation mode.")
    run_id = run_id or uuid4().hex
    with database() as connection:
        cutoff = (datetime.now(timezone.utc) - timedelta(hours=RUN_RETENTION_HOURS)).isoformat()
        connection.execute("DELETE FROM financial_movements WHERE run_id IN (SELECT id FROM demo_runs WHERE created_at < ?)", (cutoff,))
        connection.execute("DELETE FROM demo_runs WHERE created_at < ?", (cutoff,))
        connection.execute(
            "INSERT INTO demo_runs(id,project_name,payload_json,mode,owner_id,created_at) VALUES(?,?,?,?,?,?)",
            (run_id, project_name[:120], json.dumps(payload, separators=(",", ":")), mode, visitor_id(), datetime.now(timezone.utc).isoformat()),
        )
    return run_id


def get_run(run_id):
    with database() as connection:
        cutoff = (datetime.now(timezone.utc) - timedelta(hours=RUN_RETENTION_HOURS)).isoformat()
        connection.execute("DELETE FROM financial_movements WHERE run_id IN (SELECT id FROM demo_runs WHERE created_at < ?)", (cutoff,))
        connection.execute("DELETE FROM demo_runs WHERE created_at < ?", (cutoff,))
        row = connection.execute("SELECT * FROM demo_runs WHERE id=? AND owner_id=?", (run_id, visitor_id())).fetchone()
    if not row:
        return None
    payload = json.loads(row["payload_json"])
    ensure_daily_payload(payload, row["mode"])
    return {
        "id": row["id"], "project_name": row["project_name"], "mode": row["mode"],
        "created_at": row["created_at"], **payload,
    }


STATUS_LABELS = {
    "matched": "Matched",
    "mismatch": "Mismatch",
    "missing_a": "Missing in System A",
    "missing_b": "Missing in System B",
    "duplicate": "Duplicate detected",
}


def ensure_daily_payload(payload, mode):
    if "source_rows_a" not in payload or "source_rows_b" not in payload:
        payload["source_rows_a"] = [dict(row["system_a"]) for row in payload["results"] if row.get("system_a")]
        payload["source_rows_b"] = [dict(row["system_b"]) for row in payload["results"] if row.get("system_b")]
    if "daily_results" not in payload:
        payload["daily_results"] = daily_matches(
            payload["source_rows_a"], payload["source_rows_b"], detailed=mode == "detailed"
        )


def recalculate_payload(payload, mode):
    """Rebuild row and project totals after an audited amount adjustment."""
    for row in payload["results"]:
        left, right = row.get("system_a"), row.get("system_b")
        row["amount_variance"] = round((left or {}).get("amount", 0) - (right or {}).get("amount", 0), 2)
        if left and right:
            differences = [item for item in row.get("differences", []) if item != "Amount"]
            if abs(row["amount_variance"]) >= 0.01:
                differences.append("Amount")
            row["differences"] = differences
            duplicate = left.get("row_count", 1) > 1 or right.get("row_count", 1) > 1
            row["status"] = "duplicate" if duplicate else ("mismatch" if differences else "matched")
            row["status_label"] = STATUS_LABELS[row["status"]]

    counts = Counter(row["status"] for row in payload["results"])
    summary = payload["summary"]
    total = len(payload["results"])
    summary.update({
        "total": total,
        "matched": counts["matched"],
        "mismatch": counts["mismatch"],
        "missing_a": counts["missing_a"],
        "missing_b": counts["missing_b"],
        "duplicates": counts["duplicate"],
        "match_rate": round(counts["matched"] * 100 / total, 1) if total else 0,
        "amount_variance": round(sum(row["amount_variance"] for row in payload["results"]), 2),
    })
    if mode == "detailed":
        summary["pricing_variance_a"] = round(sum(
            (row.get("system_a") or {}).get("pricing_variance", 0) for row in payload["results"]
        ), 2)
        summary["pricing_variance_b"] = round(sum(
            (row.get("system_b") or {}).get("pricing_variance", 0) for row in payload["results"]
        ), 2)
    if "source_rows_a" in payload and "source_rows_b" in payload:
        payload["daily_results"] = daily_matches(
            payload["source_rows_a"], payload["source_rows_b"], detailed=mode == "detailed"
        )
    return payload


def update_source_snapshot(payload, side, reference, delta, mode):
    key = "source_rows_a" if side == "system_a" else "source_rows_b"
    members = [row for row in payload.get(key, []) if row["reference"] == reference]
    remaining = round(delta, 2)
    for row in members:
        if remaining >= 0:
            change = remaining
        else:
            change = max(remaining, -row["amount"])
        if change:
            row["amount"] = round(row["amount"] + change, 2)
            if mode == "detailed":
                row["pricing_variance"] = round(row["amount"] - row.get("calculated_amount", 0), 2)
            remaining = round(remaining - change, 2)
        if not remaining:
            break
    if remaining:
        raise ValueError("The source booking amounts cannot support this adjustment.")


def list_projects(mode):
    cutoff = (datetime.now(timezone.utc) - timedelta(hours=RUN_RETENTION_HOURS)).isoformat()
    with database() as connection:
        connection.execute("DELETE FROM financial_movements WHERE run_id IN (SELECT id FROM demo_runs WHERE created_at < ?)", (cutoff,))
        connection.execute("DELETE FROM demo_runs WHERE created_at < ?", (cutoff,))
        rows = connection.execute(
            "SELECT * FROM demo_runs WHERE mode=? AND owner_id=? ORDER BY created_at DESC, project_name", (mode, visitor_id())
        ).fetchall()
        movement_totals = {
            row["run_id"]: row["total"] for row in connection.execute(
                "SELECT run_id, ROUND(SUM(movement), 2) AS total FROM financial_movements GROUP BY run_id"
            )
        }
    projects = []
    for row in rows:
        payload = json.loads(row["payload_json"])
        projects.append({
            "id": row["id"], "project_name": row["project_name"], "mode": row["mode"],
            "created_at": row["created_at"], "summary": payload["summary"],
            "movement_total": movement_totals.get(row["id"], 0),
        })
    return projects


def ensure_sample_projects(mode):
    definitions = (
        ("may", "May 2026 marine operations", 0),
        ("june", "June 2026 marine operations", 25),
        ("july", "July 2026 marine operations", -40),
    )
    with database() as connection:
        existing = {row["id"] for row in connection.execute("SELECT id FROM demo_runs WHERE mode=? AND owner_id=?", (mode, visitor_id()))}
    for month, name, delta in definitions:
        run_id = sample_project_id(mode, month)
        if run_id in existing:
            continue
        payload = reconcile(records(SYSTEM_A), records(SYSTEM_B)) if mode == "standard" else reconcile_detailed(
            detailed_records(DETAILED_A), detailed_records(DETAILED_B)
        )
        payload = json.loads(json.dumps(payload))
        if delta:
            target = next(row for row in payload["results"] if row.get("system_a") and row.get("system_b"))
            target["system_a"]["amount"] = round(target["system_a"]["amount"] + delta, 2)
            update_source_snapshot(payload, "system_a", target["reference"], delta, mode)
            if mode == "detailed":
                target["system_a"]["pricing_variance"] = round(
                    target["system_a"]["amount"] - target["system_a"].get("calculated_amount", 0), 2
                )
            recalculate_payload(payload, mode)
        store_run(name, payload, mode=mode, run_id=run_id)


def selected_projects(mode, project_ids):
    unique_ids = list(dict.fromkeys(project_ids))[:12]
    projects = []
    resolved_ids = set()
    for run_id in unique_ids:
        for month in ("may", "june", "july"):
            if run_id == f"portfolio-{mode}-{month}":
                run_id = sample_project_id(mode, month)
                break
        if run_id in resolved_ids:
            continue
        resolved_ids.add(run_id)
        run = get_run(run_id)
        if run and run["mode"] == mode:
            projects.append(run)
    return projects


def movement_rows(run_ids):
    if not run_ids:
        return []
    placeholders = ",".join("?" for _ in run_ids)
    with database() as connection:
        rows = connection.execute(
            f"""SELECT m.*, r.project_name FROM financial_movements m
                JOIN demo_runs r ON r.id=m.run_id
                WHERE m.run_id IN ({placeholders}) ORDER BY m.id DESC""",
            run_ids,
        ).fetchall()
    return [dict(row) for row in rows]


def combined_daily_rows(mode, projects):
    source_a = [dict(row, project_name=project["project_name"]) for project in projects for row in project["source_rows_a"]]
    source_b = [dict(row, project_name=project["project_name"]) for project in projects for row in project["source_rows_b"]]
    return [{**row, "project_id": "combined"} for row in daily_matches(
        source_a, source_b, detailed=mode == "detailed"
    )]


def combined_workspace_payload(mode, projects):
    rows = []
    for project in projects:
        for result in project["results"]:
            rows.append({**result, "project_id": project["id"], "project_name": project["project_name"]})
    daily_rows = combined_daily_rows(mode, projects)
    movements = movement_rows([project["id"] for project in projects])
    return {
        "mode": mode,
        "projects": projects,
        "rows": rows,
        "daily_rows": daily_rows,
        "movements": movements,
        "summary": {
            "projects": len(projects),
            "references": len(rows),
            "matched": sum(row["status"] == "matched" for row in rows),
            "amount_variance": round(sum(row["amount_variance"] for row in rows), 2),
            "movement_total": round(sum(row["movement"] for row in movements), 2),
        },
    }


def validate_file(upload):
    if not upload or not upload.filename:
        raise ValueError("Select both System A and System B files.")
    suffix = Path(upload.filename).suffix.lower()
    if suffix not in {".xlsx", ".csv"}:
        raise ValueError("Only XLSX and CSV files are accepted.")
    upload.stream.seek(0, os.SEEK_END)
    size = upload.stream.tell()
    upload.stream.seek(0)
    if size <= 0:
        raise ValueError(f"{upload.filename}: the file is empty.")
    if size > MAX_UPLOAD_BYTES:
        raise ValueError(f"{upload.filename}: the file exceeds the 5 MB demo limit.")
    if suffix == ".xlsx":
        if not zipfile.is_zipfile(upload.stream):
            upload.stream.seek(0)
            raise ValueError(f"{upload.filename}: invalid Excel file.")
        upload.stream.seek(0)
        with zipfile.ZipFile(upload.stream) as archive:
            uncompressed_size = sum(item.file_size for item in archive.infolist())
            if uncompressed_size > MAX_UNCOMPRESSED_BYTES:
                raise ValueError(f"{upload.filename}: expanded workbook is too large for the public demo.")
    upload.stream.seek(0)


def excel_safe(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def format_export_sheet(sheet, widths):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173E47")
        cell.alignment = Alignment(horizontal="center")
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def add_daily_sheet(workbook, rows, mode, include_project=False):
    sheet = workbook.create_sheet("Daily Matching")
    headers = ["Project"] if include_project else []
    headers.extend(["Date", "Company"])
    if mode == "detailed":
        headers.append("Trip Type")
    headers.extend(["Status", "Rows A", "Rows B", "Adults A", "Adults B", "Children A", "Children B"])
    if mode == "detailed":
        headers.extend(["Pickup A", "Pickup B", "Transfer Rows A", "Transfer Rows B", "Chargeable A", "Chargeable B", "Adult Rate A", "Adult Rate B", "Calculated A", "Calculated B", "Calculated Variance"])
    headers.extend(["Amount A", "Amount B", "Variance"])
    if mode == "detailed":
        headers.extend(["Pricing Variance A", "Pricing Variance B"])
    headers.extend(["References A", "References B"])
    sheet.append(headers)
    for row in rows:
        values = [excel_safe(row["project_name"])] if include_project else []
        values.extend([row["date"], excel_safe(row["company"])])
        if mode == "detailed":
            values.append(excel_safe(row["trip"]))
        values.extend([row["status_label"], row["rows_a"], row["rows_b"], row["adult_a"], row["adult_b"], row["child_a"], row["child_b"]])
        if mode == "detailed":
            values.extend([excel_safe(row["pickup_a"]), excel_safe(row["pickup_b"]), row["transfer_a"], row["transfer_b"], row["chargeable_a"], row["chargeable_b"], row["rate_a"], row["rate_b"], row["calculated_a"], row["calculated_b"], row["calculated_variance"]])
        values.extend([row["amount_a"], row["amount_b"], row["variance"]])
        if mode == "detailed":
            values.extend([row["pricing_variance_a"], row["pricing_variance_b"]])
        values.extend([excel_safe(row["references_a"]), excel_safe(row["references_b"])])
        sheet.append(values)
    widths = ([28] if include_project else []) + [16, 28] + ([26] if mode == "detailed" else []) + [30, 12, 12, 14, 14, 14, 14]
    if mode == "detailed":
        widths.extend([28, 28, 12, 12, 16, 16, 16, 16, 18, 18, 18])
    widths.extend([18, 18, 18])
    if mode == "detailed":
        widths.extend([18, 18])
    format_export_sheet(sheet, widths + [28, 28])


def csrf_token():
    if "csrf_token" not in session:
        session["csrf_token"] = secrets.token_urlsafe(24)
    return session["csrf_token"]


@app.before_request
def protect_posts():
    if request.method == "POST":
        expected = session.get("csrf_token")
        submitted = request.form.get("csrf_token")
        if not expected or not submitted or not secrets.compare_digest(expected, submitted):
            abort(400, description="The form expired. Refresh the page and try again.")


@app.after_request
def security_headers(response):
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Content-Security-Policy"] = (
        "default-src 'self'; style-src 'self'; script-src 'self'; "
        "img-src 'self' data:; object-src 'none'; base-uri 'self'; form-action 'self'"
    )
    if request.path.startswith("/results") or request.method == "POST":
        response.headers["Cache-Control"] = "no-store"
    return response


app.jinja_env.globals["csrf_token"] = csrf_token
init_database()


SUPPORTED_LANGUAGES = {"en", "ar"}


def current_language():
    language = session.get("language", "en")
    return language if language in SUPPORTED_LANGUAGES else "en"


@app.context_processor
def inject_language():
    language = current_language()
    return {"language": language, "text_direction": "rtl" if language == "ar" else "ltr"}


@app.get("/language/<language>")
def set_language(language):
    if language not in SUPPORTED_LANGUAGES:
        abort(404)
    next_path = request.args.get("next", "/")
    if not next_path.startswith("/") or next_path.startswith("//") or "\\" in next_path:
        next_path = "/"
    session["language"] = language
    return redirect(next_path)


@app.get("/")
def hub():
    return render_template("hub.html")


@app.get("/try/<mode>")
def try_live_demo(mode):
    if mode not in {"standard", "detailed"}:
        abort(404)
    ensure_sample_projects(mode)
    return redirect(url_for(
        "combined_projects", mode=mode,
        project=[sample_project_id(mode, "may"), sample_project_id(mode, "june")],
    ))


@app.get("/standard")
def index():
    ensure_sample_projects("standard")
    return render_template("index.html", projects=list_projects("standard"))


@app.post("/demo")
def load_demo():
    payload = reconcile(records(SYSTEM_A), records(SYSTEM_B))
    run_id = store_run("Synthetic marine-tour operations demo", payload, mode="standard")
    return redirect(url_for("results", run_id=run_id))


@app.post("/reconcile")
def upload_and_reconcile():
    try:
        system_a_file = request.files.get("system_a")
        system_b_file = request.files.get("system_b")
        validate_file(system_a_file)
        validate_file(system_b_file)
        system_a = read_dataset(system_a_file.stream, system_a_file.filename)
        system_b = read_dataset(system_b_file.stream, system_b_file.filename)
        payload = reconcile(system_a, system_b)
        run_id = store_run(request.form.get("project_name", "").strip() or "Reconciliation review", payload, mode="standard")
        return redirect(url_for("results", run_id=run_id))
    except ValueError as exc:
        flash(str(exc), "error")
        return render_template("index.html"), 400


@app.get("/results/<run_id>")
def results(run_id):
    run = get_run(run_id)
    if not run:
        abort(404)
    return render_template("results.html", run=run)


@app.get("/results/<run_id>/export.xlsx")
def export_results(run_id):
    run = get_run(run_id)
    if not run:
        abort(404)
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet.append(["Metric", "Value"])
    summary_rows = [
        ("Review name", excel_safe(run["project_name"])),
        ("Created UTC", run["created_at"]),
        ("Total references", run["summary"]["total"]),
        ("Exact matches", run["summary"]["matched"]),
        ("Field-level mismatches", run["summary"]["mismatch"]),
        ("Missing in System A", run["summary"]["missing_a"]),
        ("Missing in System B", run["summary"]["missing_b"]),
        ("Duplicate references", run["summary"]["duplicates"]),
        ("Exact match rate", run["summary"]["match_rate"] / 100),
        ("System A source rows", run["summary"]["system_a_rows"]),
        ("System B source rows", run["summary"]["system_b_rows"]),
        ("Net amount variance", run["summary"]["amount_variance"]),
    ]
    for row in summary_rows:
        summary_sheet.append(row)
    summary_sheet["B10"].number_format = "0.0%"
    format_export_sheet(summary_sheet, [30, 32])

    sheet = workbook.create_sheet("Reconciliation Results")
    headers = [
        "Reference", "Status", "Differences", "System A Date", "System B Date",
        "System A Company", "System B Company", "System A Service", "System B Service",
        "A Adult", "B Adult", "A Child", "B Child", "A Infant", "B Infant",
        "System A Amount", "System B Amount", "Variance", "A Source Rows", "B Source Rows",
    ]
    sheet.append(headers)
    for result in run["results"]:
        left = result["system_a"] or {}
        right = result["system_b"] or {}
        sheet.append([
            excel_safe(result["reference"]), result["status_label"], ", ".join(result["differences"]) or "None",
            left.get("date", ""), right.get("date", ""),
            excel_safe(left.get("company", "")), excel_safe(right.get("company", "")),
            excel_safe(left.get("service", "")), excel_safe(right.get("service", "")),
            left.get("adult", 0), right.get("adult", 0), left.get("child", 0), right.get("child", 0),
            left.get("infant", 0), right.get("infant", 0), left.get("amount", 0), right.get("amount", 0),
            result["amount_variance"], left.get("row_count", 0), right.get("row_count", 0),
        ])
    format_export_sheet(sheet, [18, 23, 34, 16, 16, 26, 26, 24, 24] + [12] * 11)
    add_daily_sheet(workbook, run["daily_results"], "standard")

    mapping_sheet = workbook.create_sheet("Company Mappings")
    mapping_sheet.append(["System A Label", "System B Label", "Normalized A", "Normalized B", "Result"])
    for mapping in run["mappings"]:
        mapping_sheet.append([
            excel_safe(mapping["system_a"]), excel_safe(mapping["system_b"]),
            excel_safe(mapping["system_a_canonical"]), excel_safe(mapping["system_b_canonical"]),
            "Aligned" if mapping["mapped"] else "Review required",
        ])
    format_export_sheet(mapping_sheet, [28, 28, 28, 28, 20])
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name="reconcileflow-demo-results.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.get("/sample-data/<system>.xlsx")
def sample_workbook(system):
    rows = SYSTEM_A if system.lower() == "a" else SYSTEM_B if system.lower() == "b" else None
    if rows is None:
        abort(404)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"System {system.upper()}"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    format_export_sheet(sheet, [18, 15, 28, 24, 12, 12, 12, 16])
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=f"system-{system.lower()}-demo.xlsx")


@app.get("/detailed")
def detailed_workspace():
    ensure_sample_projects("detailed")
    return render_template("detailed.html", projects=list_projects("detailed"))


@app.post("/detailed/demo")
def load_detailed_demo():
    payload = reconcile_detailed(detailed_records(DETAILED_A), detailed_records(DETAILED_B))
    run_id = store_run("Synthetic detailed marine-tour review", payload, mode="detailed")
    return redirect(url_for("detailed_results", run_id=run_id))


@app.post("/detailed/reconcile")
def upload_detailed():
    try:
        system_a_file = request.files.get("system_a")
        system_b_file = request.files.get("system_b")
        validate_file(system_a_file)
        validate_file(system_b_file)
        left = read_detailed_dataset(system_a_file.stream, system_a_file.filename)
        right = read_detailed_dataset(system_b_file.stream, system_b_file.filename)
        run_id = store_run(
            request.form.get("project_name", "").strip() or "Detailed reconciliation review",
            reconcile_detailed(left, right),
            mode="detailed",
        )
        return redirect(url_for("detailed_results", run_id=run_id))
    except ValueError as exc:
        flash(str(exc), "error")
        return render_template("detailed.html"), 400


@app.get("/detailed/results/<run_id>")
def detailed_results(run_id):
    run = get_run(run_id)
    if not run:
        abort(404)
    return render_template("detailed_results.html", run=run)


@app.get("/detailed/results/<run_id>/export.xlsx")
def detailed_export(run_id):
    run = get_run(run_id)
    if not run:
        abort(404)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Metric", "Value"])
    for label, key in (
        ("Total references", "total"), ("Exact matches", "matched"), ("Mismatches", "mismatch"),
        ("Missing in System A", "missing_a"), ("Missing in System B", "missing_b"),
        ("Duplicates", "duplicates"), ("Match rate", "match_rate"),
        ("Net amount variance", "amount_variance"), ("System A pricing variance", "pricing_variance_a"),
        ("System B pricing variance", "pricing_variance_b"),
    ):
        summary.append([label, run["summary"][key]])
    format_export_sheet(summary, [32, 22])
    detail = workbook.create_sheet("Detailed Results")
    detail.append([
        "Reference", "Status", "Differences", "A Date", "B Date", "A Company", "B Company",
        "A Trip", "B Trip", "A Hotel or Pickup", "B Hotel or Pickup", "A Adult", "B Adult",
        "A Child", "B Child", "A Infant", "B Infant", "A Chargeable Pax", "B Chargeable Pax",
        "A Adult Rate", "B Adult Rate", "A Amount", "B Amount", "Amount Variance",
        "A Pricing Variance", "B Pricing Variance",
    ])
    for row in run["results"]:
        left, right = row["system_a"] or {}, row["system_b"] or {}
        detail.append([
            excel_safe(row["reference"]), row["status_label"], ", ".join(row["differences"]) or "None",
            left.get("date", ""), right.get("date", ""), excel_safe(left.get("company", "")),
            excel_safe(right.get("company", "")), excel_safe(left.get("service", "")),
            excel_safe(right.get("service", "")), excel_safe(left.get("pickup", "")),
            excel_safe(right.get("pickup", "")), left.get("adult", 0), right.get("adult", 0),
            left.get("child", 0), right.get("child", 0), left.get("infant", 0), right.get("infant", 0),
            left.get("chargeable_pax", 0), right.get("chargeable_pax", 0), left.get("adult_rate", 0),
            right.get("adult_rate", 0), left.get("amount", 0), right.get("amount", 0), row["amount_variance"],
            left.get("pricing_variance", 0), right.get("pricing_variance", 0),
        ])
    format_export_sheet(detail, [18, 22, 34] + [16, 16, 26, 26, 22, 22, 24, 24] + [13] * 15)
    add_daily_sheet(workbook, run["daily_results"], "detailed")
    reference = workbook.create_sheet("Reference Matching")
    reference.append(["System A Date", "System B Date", "Reference", "System A Amount", "System B Amount", "Variance", "Status"])
    for row in run["results"]:
        left, right = row["system_a"] or {}, row["system_b"] or {}
        reference.append([left.get("date", ""), right.get("date", ""), excel_safe(row["reference"]), left.get("amount", 0), right.get("amount", 0), row["amount_variance"], row["status_label"]])
    format_export_sheet(reference, [18, 18, 20, 18, 18, 18, 24])
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="reconcileflow-detailed-demo.xlsx")


@app.get("/detailed/sample-data/<system>.xlsx")
def detailed_sample(system):
    rows = DETAILED_A if system.lower() == "a" else DETAILED_B if system.lower() == "b" else None
    if rows is None:
        abort(404)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"Detailed {system.upper()}"
    sheet.append(DETAILED_HEADERS)
    for row in rows:
        sheet.append(row)
    format_export_sheet(sheet, [18, 15, 28, 24, 26, 12, 12, 12, 16, 16])
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=f"detailed-system-{system.lower()}-demo.xlsx")


@app.get("/projects/<mode>/combined")
def combined_projects(mode):
    if mode not in {"standard", "detailed"}:
        abort(404)
    ensure_sample_projects(mode)
    projects = selected_projects(mode, request.args.getlist("project"))
    if not projects:
        return redirect(url_for("index" if mode == "standard" else "detailed_workspace"))
    return render_template("project_workspace.html", workspace=combined_workspace_payload(mode, projects))


@app.get("/projects/<mode>/combined/export.xlsx")
def combined_projects_export(mode):
    if mode not in {"standard", "detailed"}:
        abort(404)
    ensure_sample_projects(mode)
    projects = selected_projects(mode, request.args.getlist("project"))
    if not projects:
        abort(404)
    data = combined_workspace_payload(mode, projects)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Combined Summary"
    summary.append(["Metric", "Value"])
    for label, key in (
        ("Projects", "projects"), ("References", "references"), ("Exact matches", "matched"),
        ("Amount variance", "amount_variance"), ("Settlement movement", "movement_total"),
    ):
        summary.append([label, data["summary"][key]])
    summary.append([])
    summary.append(["Included projects", "Mode"])
    for project in projects:
        summary.append([excel_safe(project["project_name"]), mode.title()])
    format_export_sheet(summary, [34, 24])

    bookings = workbook.create_sheet("Combined Bookings")
    headers = ["Project", "Reference", "Status", "Differences", "System A Company", "System B Company"]
    if mode == "detailed":
        headers.extend(["System A Trip", "System B Trip", "System A Hotel or Pickup", "System B Hotel or Pickup"])
    headers.extend(["System A Amount", "System B Amount", "Variance"])
    bookings.append(headers)
    for row in data["rows"]:
        left, right = row.get("system_a") or {}, row.get("system_b") or {}
        values = [
            excel_safe(row["project_name"]), excel_safe(row["reference"]), row["status_label"],
            ", ".join(row["differences"]) or "None", excel_safe(left.get("company", "")),
            excel_safe(right.get("company", "")),
        ]
        if mode == "detailed":
            values.extend([
                excel_safe(left.get("service", "")), excel_safe(right.get("service", "")),
                excel_safe(left.get("pickup", "")), excel_safe(right.get("pickup", "")),
            ])
        values.extend([left.get("amount", 0), right.get("amount", 0), row["amount_variance"]])
        bookings.append(values)
    format_export_sheet(bookings, [28, 18, 22, 34, 26, 26] + ([22, 22, 26, 26] if mode == "detailed" else []) + [18, 18, 18])

    add_daily_sheet(workbook, data["daily_rows"], mode, include_project=True)

    ledger = workbook.create_sheet("Financial Movements")
    ledger.append(["Adjustment UTC", "Project", "Reference", "Source", "Before", "After", "Movement", "Action"])
    for row in data["movements"]:
        ledger.append([
            row["occurred_at"], excel_safe(row["project_name"]), excel_safe(row["reference"]),
            "System A" if row["side"] == "system_a" else "System B", row["before_amount"],
            row["after_amount"], row["movement"], row["action"],
        ])
    format_export_sheet(ledger, [24, 28, 18, 14, 16, 16, 16, 28])
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output, as_attachment=True, download_name=f"reconcileflow-{mode}-combined-projects.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.post("/projects/<run_id>/adjustment")
def adjust_project_amount(run_id):
    reference = request.form.get("reference", "").strip()
    side = request.form.get("side", "").strip().lower()
    if side not in {"system_a", "system_b"} or not reference:
        return jsonify({"ok": False, "error": "Invalid booking adjustment."}), 400
    try:
        after_amount = round(float(request.form.get("amount", "")), 2)
    except (TypeError, ValueError):
        return jsonify({"ok": False, "error": "Amount must be numeric."}), 400
    if not math.isfinite(after_amount) or after_amount < 0 or after_amount > 1_000_000_000:
        return jsonify({"ok": False, "error": "Amount is outside the accepted range."}), 400

    occurred_at = datetime.now(timezone.utc).isoformat()
    with database() as connection:
        connection.execute("BEGIN IMMEDIATE")
        record = connection.execute("SELECT * FROM demo_runs WHERE id=? AND owner_id=?", (run_id, visitor_id())).fetchone()
        if not record:
            abort(404)
        payload = json.loads(record["payload_json"])
        ensure_daily_payload(payload, record["mode"])
        result = next((row for row in payload["results"] if row["reference"] == reference), None)
        source = result.get(side) if result else None
        if source is None:
            return jsonify({"ok": False, "error": "The selected source has no booking to adjust."}), 400
        before_amount = round(float(source.get("amount", 0)), 2)
        if before_amount == after_amount:
            return jsonify({"ok": False, "error": "Enter an amount different from the current value."}), 400
        source["amount"] = after_amount
        update_source_snapshot(payload, side, reference, after_amount - before_amount, record["mode"])
        if record["mode"] == "detailed":
            source["pricing_variance"] = round(after_amount - float(source.get("calculated_amount", 0)), 2)
        recalculate_payload(payload, record["mode"])
        movement = round(before_amount - after_amount, 2)
        action = f"{'System A' if side == 'system_a' else 'System B'} amount adjusted"
        connection.execute(
            "UPDATE demo_runs SET payload_json=? WHERE id=?",
            (json.dumps(payload, separators=(",", ":")), run_id),
        )
        cursor = connection.execute(
            """INSERT INTO financial_movements(
                   run_id,mode,occurred_at,reference,side,before_amount,after_amount,movement,action
               ) VALUES(?,?,?,?,?,?,?,?,?)""",
            (run_id, record["mode"], occurred_at, reference, side, before_amount, after_amount, movement, action),
        )

    return jsonify({
        "ok": True,
        "movement": {
            "id": cursor.lastrowid, "occurred_at": occurred_at, "reference": reference,
            "side": side, "before": before_amount, "after": after_amount,
            "movement": movement, "action": action, "project": record["project_name"],
        },
        "row": {
            "status": result["status"], "status_label": result["status_label"],
            "differences": result["differences"], "amount_variance": result["amount_variance"],
        },
        "daily_html": render_template(
            "_daily_match_rows.html",
            rows=combined_daily_rows(
                record["mode"],
                selected_projects(record["mode"], request.form.getlist("project") or [run_id]),
            ),
            is_detailed=record["mode"] == "detailed",
        ),
        "summary": payload["summary"],
    })


@app.get("/invoice-tracking")
def invoice_tracking():
    tab = request.args.get("tab", "overview")
    if tab not in {"overview", "companies", "names", "analysis", "activity"}:
        tab = "overview"
    return render_template("invoice_tracking.html", data=invoice_demo_payload(), tab=tab)


def invoice_update_response(code, changes, action):
    ensure_invoice_demo()
    owner = visitor_id()
    with database() as connection:
        connection.execute("BEGIN IMMEDIATE")
        row = connection.execute(
            "SELECT * FROM invoice_demo_companies WHERE owner_id=? AND code=?", (owner, code)
        ).fetchone()
        if not row:
            abort(404)
        old = dict(row)
        if all(old[key] == value for key, value in changes.items()):
            return jsonify({"ok": False, "error": "No changes to save."}), 400
        before = old["invoice_minor"] - old["paid_minor"] - old["payment_minor"]
        after = changes.get("invoice_minor", old["invoice_minor"]) - changes.get("paid_minor", old["paid_minor"]) - changes.get("payment_minor", old["payment_minor"])
        assignments = ", ".join(f"{key}=?" for key in changes)
        connection.execute(
            f"UPDATE invoice_demo_companies SET {assignments} WHERE owner_id=? AND code=?",
            (*changes.values(), owner, code),
        )
        connection.execute(
            """INSERT INTO invoice_demo_events(
                   owner_id,company_code,company_name,occurred_at,action,before_minor,after_minor
               ) VALUES(?,?,?,?,?,?,?)""",
            (owner, code, changes.get("name", old["name"]), datetime.now(timezone.utc).isoformat(), action, before, after),
        )
    data = invoice_demo_payload()
    company = next(row for row in data["companies"] if row["code"] == code)
    return jsonify({"ok": True, "company": company, "summary": data["summary"],
                    "row_html": render_template("_invoice_demo_name_row.html", row=company)})


@app.post("/invoice-tracking/company/<code>/tracking")
def invoice_tracking_update(code):
    stage = request.form.get("stage", "")
    if stage not in STAGES:
        return jsonify({"ok": False, "error": "Choose a valid stage."}), 400
    try:
        payment = amount_to_minor(request.form.get("payment", ""))
    except ValueError as error:
        return jsonify({"ok": False, "error": str(error)}), 400
    return invoice_update_response(code, {"payment_minor": payment, "stage": stage}, "Payment or stage updated")


@app.post("/invoice-tracking/company/<code>/identity")
def invoice_identity_update(code):
    name = " ".join(request.form.get("name", "").split())[:120]
    if not name:
        return jsonify({"ok": False, "error": "Company name is required."}), 400
    try:
        invoice = amount_to_minor(request.form.get("invoice", ""))
        paid = amount_to_minor(request.form.get("paid", ""))
    except ValueError as error:
        return jsonify({"ok": False, "error": str(error)}), 400
    notes = request.form.get("notes", "").strip()[:1000]
    return invoice_update_response(
        code, {"name": name, "invoice_minor": invoice, "paid_minor": paid, "notes": notes},
        "Company identity or opening amounts updated",
    )


@app.post("/invoice-tracking/companies")
def invoice_add_company():
    name = " ".join(request.form.get("name", "").split())[:120]
    if not name:
        return jsonify({"ok": False, "error": "Company name is required."}), 400
    try:
        invoice = amount_to_minor(request.form.get("invoice", "0"))
        paid = amount_to_minor(request.form.get("paid", "0"))
    except ValueError as error:
        return jsonify({"ok": False, "error": str(error)}), 400
    ensure_invoice_demo()
    owner = visitor_id()
    with database() as connection:
        connection.execute("BEGIN IMMEDIATE")
        codes = [row["code"] for row in connection.execute(
            "SELECT code FROM invoice_demo_companies WHERE owner_id=?", (owner,)
        )]
        number = max((int(code[4:]) for code in codes if code.startswith("CMP-") and code[4:].isdigit()), default=0) + 1
        code = f"CMP-{number:03d}"
        connection.execute(
            """INSERT INTO invoice_demo_companies(
                   owner_id,code,name,invoice_minor,paid_minor,payment_minor,stage,notes,created_at
               ) VALUES(?,?,?,?,?,?,?,?,?)""",
            (owner, code, name, invoice, paid, 0, "waiting", "", datetime.now(timezone.utc).isoformat()),
        )
        connection.execute(
            """INSERT INTO invoice_demo_events(
                   owner_id,company_code,company_name,occurred_at,action,before_minor,after_minor
               ) VALUES(?,?,?,?,?,?,?)""",
            (owner, code, name, datetime.now(timezone.utc).isoformat(), "Company added", 0, invoice - paid),
        )
    data = invoice_demo_payload()
    company = next(row for row in data["companies"] if row["code"] == code)
    return jsonify({"ok": True, "company": company, "summary": data["summary"],
                    "row_html": render_template("_invoice_demo_name_row.html", row=company)})


@app.get("/invoice-tracking/export.xlsx")
def invoice_tracking_export():
    data = invoice_demo_payload(request.args.get("q", ""), request.args.get("direction", ""))
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Metric", "Value"])
    for label, key in (
        ("Companies", "companies"), ("Invoice value", "invoice"), ("Collected", "collected"),
        ("Receivable", "receivable"), ("Payable", "payable"), ("Net balance", "net"),
        ("Collection rate", "collection_rate"), ("Completed cycles", "completed"),
    ):
        summary.append([label, data["summary"][key]])
    format_export_sheet(summary, [28, 20])
    companies = workbook.create_sheet("Company Tracking")
    companies.append(["Code", "Company", "Stage", "Invoice", "Paid", "Payment", "Balance", "Direction", "Progress", "Notes"])
    for row in data["companies"]:
        companies.append([row["code"], excel_safe(row["name"]), row["stage"].title(), row["invoice"], row["paid"], row["payment"], row["balance"], row["direction"].title(), row["progress"] / 100, excel_safe(row["notes"])])
    format_export_sheet(companies, [14, 28, 16, 16, 16, 16, 16, 16, 14, 34])
    for cell in companies["I"][1:]:
        cell.number_format = "0%"
    activity = workbook.create_sheet("Activity Log")
    activity.append(["Timestamp", "Company", "Action", "Balance Before", "Balance After", "Movement"])
    for row in data["activities"]:
        activity.append([row["time"], row["company"], row["action"], row["before"], row["after"], row["after"] - row["before"]])
    format_export_sheet(activity, [22, 28, 30, 18, 18, 18])
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="reconcileflow-invoice-tracking-demo.xlsx")


@app.get("/control-center")
def control_center():
    tab = request.args.get("tab", "overview")
    if tab not in {"overview", "companies", "settlements", "no-shows", "delivery"}:
        tab = "overview"
    return render_template("control_center.html", data=control_center_payload(), tab=tab)


@app.get("/control-center/export.xlsx")
def control_center_export():
    data = control_center_payload()
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Executive Summary"
    summary.append(["Metric", "Value"])
    for label, key in (
        ("Combined projects", "projects"), ("Portfolio companies", "companies"),
        ("Completion progress", "progress"), ("Receivable variance", "receivable"),
        ("Payable variance", "payable"), ("Net variance", "net"),
        ("Settlement movement", "settlement_movement"), ("No-show bookings", "no_show_bookings"),
        ("No-show value", "no_show_value"), ("Resolved exceptions", "resolved"),
        ("Open exceptions", "open"), ("Delivered companies", "delivered"),
    ):
        summary.append([label, data["summary"][key]])
    format_export_sheet(summary, [30, 22])

    projects = workbook.create_sheet("Projects")
    projects.append(["Project", "Period", "Progress", "Companies", "Open", "Receivable", "Payable"])
    for row in data["projects"]:
        projects.append([row["name"], row["period"], row["progress"] / 100, row["companies"], row["open"], row["receivable"], row["payable"]])
    format_export_sheet(projects, [24, 16, 14, 14, 12, 18, 18])
    for cell in projects["C"][1:]:
        cell.number_format = "0%"

    links = workbook.create_sheet("Unified Companies")
    links.append(["Unified Company", "System A Labels", "System B Labels", "Projects", "Aliases", "Confidence", "Status"])
    for row in data["company_links"]:
        links.append([excel_safe(row["canonical"]), excel_safe(row["system_a"]), excel_safe(row["system_b"]), row["projects"], row["aliases"], row["confidence"] / 100, row["status"].title()])
    format_export_sheet(links, [28, 38, 38, 12, 12, 14, 16])
    for cell in links["F"][1:]:
        cell.number_format = "0%"

    ledger = workbook.create_sheet("Settlement Ledger")
    ledger.append(["Date", "Time", "Project", "Reference", "Company", "Action", "Before", "Movement", "After", "Type"])
    for row in data["settlements"]:
        ledger.append([row["date"], row["time"], row["project"], row["reference"], excel_safe(row["company"]), row["action"], row["before"], row["movement"], row["after"], row["status"].title()])
    format_export_sheet(ledger, [16, 12, 24, 16, 28, 30, 16, 16, 16, 16])

    no_shows = workbook.create_sheet("No Shows")
    no_shows.append(["Reference", "Date", "Company", "Detected Shape", "Adult", "Child", "Infant", "Value", "Resolution", "Status"])
    for row in data["no_shows"]:
        no_shows.append([row["reference"], row["date"], excel_safe(row["company"]), row["shape"], row["adult"], row["child"], row["infant"], row["value"], row["resolution"].title(), row["status"].title()])
    format_export_sheet(no_shows, [16, 16, 28, 32, 12, 12, 12, 16, 16, 16])

    deliveries = workbook.create_sheet("Delivery Control")
    deliveries.append(["Company", "Project", "Bookings", "Approved Invoice", "Delivered On", "Status"])
    for row in data["deliveries"]:
        deliveries.append([excel_safe(row["company"]), row["project"], row["bookings"], row["invoice"], row["delivered_on"], row["status"].title()])
    format_export_sheet(deliveries, [28, 24, 14, 20, 16, 16])

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="reconcileflow-operations-control-demo.xlsx")


@app.get("/report-workshop")
def report_workshop():
    return render_template(
        "report_workshop.html",
        data=invoice_demo_payload(request.args.get("q", ""), request.args.get("direction", "")),
        report_date=datetime.now().date().isoformat(),
    )


@app.get("/health")
def health():
    return jsonify({"ok": True, "service": "ReconcileFlow"})


@app.errorhandler(400)
def bad_request(error):
    message = getattr(error, "description", "Check the submitted data and try again.")
    if request.accept_mimetypes.best == "application/json":
        return jsonify({"ok": False, "error": message}), 400
    return render_template("error.html", title="Request could not be processed", message=message), 400


@app.errorhandler(404)
def not_found(_error):
    message = "The requested demo result is unavailable or has expired."
    if request.accept_mimetypes.best == "application/json":
        return jsonify({"ok": False, "error": message}), 404
    return render_template("error.html", title="Page not found", message=message), 404


@app.errorhandler(413)
def too_large(_error):
    return render_template("error.html", title="File is too large", message="The public demo accepts files up to 5 MB."), 413


if __name__ == "__main__":
    app.run(host=os.environ.get("HOST", "127.0.0.1"), port=int(os.environ.get("PORT", "8000")), debug=False)
